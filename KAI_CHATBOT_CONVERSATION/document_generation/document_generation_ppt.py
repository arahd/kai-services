import re
import os
import json
import base64
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn,nsdecls
from pptx.oxml import parse_xml
from logger import logger,get_caller_frame
import dotenv
from urllib.parse import urlparse

from agent_framework import Agent,Message,Content

from simple_agent_pool import SimpleAgentPool

dotenv.load_dotenv()

PROJECT_NAME = "KAI"
SERVICE_NAME = "KAI_CHATBOT_CONVERSATIONAL_SERVICE"

def load_prompt_from_file(prompt_file_name:str):
    """Load prompt template from external file"""
    extras = {
        'servicename': SERVICE_NAME,
        'functionname': 'load_prompt_from_file',
        'projectname': PROJECT_NAME,
        'bot_name': 'kumaran'
    }
    try:
        # Get the directory where this script is located
        script_dir = os.path.dirname(os.path.abspath(__file__))
        prompts_dir = os.path.join(script_dir, 'prompts')
        prompt_file = os.path.join(prompts_dir, f"{prompt_file_name}.txt")
        
        # Check if file exists
        if not os.path.exists(prompt_file):
            logger.error(f"Prompt file not found: {prompt_file}", extra=extras)
            raise FileNotFoundError(f"Prompt file not found: {prompt_file}")
        
        # Read the prompt from file
        with open(prompt_file, 'r', encoding='utf-8') as f:
            prompt = f.read()
        
        logger.info(f"Successfully loaded prompt from: {prompt_file}", extra=extras)
        return prompt
        
    except Exception as e:
        logger.error(f"Error loading prompt file {prompt_file_name}: {str(e)}", extra=extras)
        raise

class TokenCounter:
    def __init__(self):
        self.total_input_tokens = 0
        self.total_output_tokens = 0
        self.total_tokens = 0

    def update_usuage(self, response):
        """Updates the token usage with new values."""
        try:
            x = response.metadata.get("metadata")[0]
            usage = x.get("usage")
            output_token_count = getattr(usage, 'completion_tokens', None) or getattr(usage, 'output_tokens', None) or 0
            input_token_count  = getattr(usage, 'prompt_tokens',     None) or getattr(usage, 'input_tokens',  None) or 0
            self.total_input_tokens+=input_token_count
            self.total_output_tokens+=output_token_count
        except Exception as e:
            pass


    def get_total_usuage(self) -> dict:
        """Returns the total accumulated token usage."""
        return {
            "prompt_tokens": self.total_input_tokens,
            "completion_tokens": self.total_output_tokens,
            "total_tokens": self.total_tokens
        }
    

async def ppt_outline_generation(model_name:str,user_query:str,token_counter:TokenCounter,simple_agent_pool:SimpleAgentPool):

    agent = simple_agent_pool.agent_pool.get(model_name,None)
    if not agent:
        agent = simple_agent_pool.agent_pool.get("gpt-4o-mini")
    
    prompt = load_prompt_from_file(prompt_file_name="pptx_document_generation")

    system_message = Message( role="system",contents=[
        Content.from_text(text=prompt)
                                               ],)
                
    user_message= Message( role="user",contents=[
                                 Content.from_text(text=user_query)
                                               ],)

    response = await agent.run(messages=[system_message,user_message])

    token_counter.update_usuage(response=response)
    response=str(response)

    if response.startswith("```") or response.endswith("```"):
                response = re.sub(r"^```(?:json)?\s*", "", response,flags=re.IGNORECASE)
                response = re.sub(r"\s*```$", "", response)
    response = response.strip()

    response = json.loads(response)
    return response


async def generate_ppt_file(user_query,model_name,simple_agent_pool:SimpleAgentPool, web_search_context: str = ''):
    extras = {
        'servicename': SERVICE_NAME,
        'functionname': get_caller_frame().f_code.co_name,
        'projectname': PROJECT_NAME,
        'bot_name':'kumaran'
    }
    token_counter = TokenCounter()
    ppt_file_bytes = None
    ppt_outline_json = None
    slide_contents = list()
    try:
        if web_search_context:
            user_query += f"\n\nContext from Web Search:\n{web_search_context}"
        # prev_slide_details = dict()
        logger.info(f"OUTLINE GENERATION FOR PPT STARTED",extra=extras)
        ppt_outline_json = await ppt_outline_generation(model_name=model_name,user_query=user_query,token_counter=token_counter,simple_agent_pool=simple_agent_pool)
        logger.info(f"Completed OUTLINE GENERATION FOR PPT",extra = extras)
        logger.info(f"Compiling PPT from generated content",extra = extras)
        ppt_file_bytes = generate_pptx_file(outline_json=ppt_outline_json)
        return {"ppt_file_bytes":ppt_file_bytes,"filename":ppt_outline_json.get("document_name"),"input_tokens":token_counter.total_input_tokens,"output_tokens":token_counter.total_output_tokens}
            
    except Exception as e:
        return "Unexpected Error During PPT generation"
    finally:
        del ppt_file_bytes
        del ppt_outline_json
        del slide_contents

LAYOUT_MAP = {
    "TITLE_SLIDE": 0,
    "TITLE_CONTENT": 1,
    "SECTION_HEADER": 2,
    "TITLE_ONLY": 5
}

def hex_to_rgb(hex_str):
    if not hex_str or not isinstance(hex_str, str): 
        return RGBColor(0, 0, 0)
    hex_str = hex_str.lstrip('#')
    if len(hex_str) == 3:
        hex_str = ''.join([c*2 for c in hex_str])
    try:
        r, g, b = tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4))
        return RGBColor(r, g, b)
    except:
        return RGBColor(0, 0, 0)

def remove_bullets(paragraph):
    """Removes bullets using parse_xml to avoid OxmlElement import issues."""
    p_pr = paragraph._p.get_or_add_pPr()
    
    # Remove any existing bullet tags via the underlying XML
    for tag in ['buChar', 'buNone', 'buAutoNum', 'buBlip', 'buFont', 'buClr', 'buSzPct']:
        child = p_pr.find(qn(f'a:{tag}'))
        if child is not None: 
            p_pr.remove(child)
    
    # Add "No Bullet" tag using parse_xml
    bu_none_xml = f'<a:buNone {nsdecls("a")}/>'
    p_pr.insert(0, parse_xml(bu_none_xml))
    
    # Reset indentation and alignment
    p_pr.set('marL', '0')
    p_pr.set('indent', '0')
    paragraph.level = 0
    paragraph.alignment = PP_ALIGN.LEFT


def calculate_text_height(text_frame, box_width_inches=9.0):
    """Calculate approximate height based on actual content added"""
    total_height = 0
    
    for para in text_frame.paragraphs:
        text = para.text
        font_size = para.font.size or Pt(18)  # Default if not set
        font_size_inches = font_size.inches if hasattr(font_size, 'inches') else font_size / 914400
        chars_per_line = int(box_width_inches / (font_size_inches * 0.6))# Estimate characters per line based on box width
        num_lines = max(1, (len(text) // chars_per_line) + 1)# Calculate number of lines this paragraph will wrap to
        
        line_height = font_size_inches * 1.2# Line height = font size * 1.2 (typical line spacing)
        
        space_after = para.space_after.inches if para.space_after else 0.05# Add paragraph spacing
        
        total_height += (num_lines * line_height) + space_after# Add paragraph spacing
    return total_height


def generate_pptx_file(outline_json,output_file="ppt.pptx"):
    prs = Presentation()

    global_bg = hex_to_rgb(outline_json.get("presentation_background_color", "#FFFFFF"))
    global_font = hex_to_rgb(outline_json.get("presentation_font_color", "#000000"))
    detailed_contents=outline_json.get("slide_content",list())

    for content in detailed_contents:
        content_height=0
        layout_name = content.get("slide_layout", "TITLE_CONTENT")
        layout_idx = LAYOUT_MAP.get(layout_name, 1) 
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        
        # Set Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = global_bg

        # Handle Title
        if slide.shapes.title:
            slide.shapes.title.text = content.get("title", "")
            slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = global_font
            

        # Handle Subtitle
        subtitle_text = content.get("subtitle", "")
        if subtitle_text:
            if layout_name in ["TITLE_SLIDE", "SECTION_HEADER"] and len(slide.placeholders) > 1:
                sp = slide.placeholders[1]
                sp.text = subtitle_text
                sp.text_frame.paragraphs[0].font.size = Pt(24)
                sp.text_frame.paragraphs[0].font.color.rgb = global_font
            elif slide.shapes.title:
                st_p = slide.shapes.title.text_frame.add_paragraph()
                st_p.text = subtitle_text
                st_p.font.size = Pt(18)
                st_p.font.italic = True
                st_p.font.color.rgb = global_font
                

        # Handle Body Text
        text_height_estimate = 0
        if layout_name == "TITLE_CONTENT" and len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = "" # Clears placeholder, leaves exactly one paragraph at index 0
            
            is_first_item = True # THIS PREVENTS THE EMPTY LINE

            # A. Paragraphs
            paragraphs = content.get("paragraph_data", [])
            for text_block in paragraphs:
                if is_first_item:
                    p = tf.paragraphs[0]
                    is_first_item = False
                else:
                    p = tf.add_paragraph()
                
                p.text = text_block
                p.font.color.rgb = global_font
                p.font.size = Pt(18)
                p.space_after = Pt(10)
                remove_bullets(p)
                text_height_estimate += 0.45

            # B. Bullets
            bullets = content.get("bullets", [])
            for b_item in bullets:
                if is_first_item:
                    p = tf.paragraphs[0]
                    is_first_item = False
                else:
                    p = tf.add_paragraph()
                
                p.text = b_item.get("text", "")
                p.level = min(b_item.get("level", 0), 2)
                p.font.color.rgb = global_font
                p.font.size = Pt(18 - (p.level * 2))
                text_height_estimate += 0.35
            content_height = calculate_text_height(tf)

        # Handle Tables
        table_data = content.get("table_data", [])
        if table_data:
            rows, cols = len(table_data), len(table_data[0])
            styles = content.get("table_styles", {})
            h_bg = hex_to_rgb(styles.get("header_bg", "#323232"))
            h_font = hex_to_rgb(styles.get("header_font", "#FFFFFF"))
            b_font = hex_to_rgb(styles.get("body_font", global_font))

            padding = Inches(0.6) if text_height_estimate > 0 else Inches(0)
            top_pos = body.top + Inches(content_height) + Inches(0.3)
            top_pos = min(top_pos, Inches(5.3))
            
            table_shape = slide.shapes.add_table(rows, cols, Inches(0.5), top_pos, Inches(9), Inches(0.4 * rows))
            table = table_shape.table
            
            for r_idx, row_data in enumerate(table_data):
                for c_idx, val in enumerate(row_data):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(val)
                    para = cell.text_frame.paragraphs[0]
                    para.font.size = Pt(12)
                    if r_idx == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = h_bg
                        para.font.color.rgb = h_font
                        para.font.bold = True
                    else:
                        para.font.color.rgb = b_font

        # Handle Footer
        # footer_text = content.get("footer", "")  # KEEPING THE CODE INCASE we need to go back to text based footers
        # if footer_text:
        #     tx_footer = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(9), Inches(0.3))
        #     fp = tx_footer.text_frame.paragraphs[0]
        #     fp.text = footer_text
        #     fp.font.size = Pt(12)
        #     fp.font.color.rgb = global_font
        #     fp.alignment = PP_ALIGN.RIGHT


        footer_image_path = "document_generation/assets/kai_logo.png"
        slide.shapes.add_picture(
            footer_image_path,
            Inches(9.35),# Left position
            Inches(6.65)# Top position 
            ) 

    pptx_buffer = BytesIO()
    prs.save(pptx_buffer)
    pptx_buffer.seek(0)
    return base64.b64encode(pptx_buffer.read()).decode("utf-8")