import os
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PyPDF2 import PdfReader, PdfWriter
import copy, json
from prompts.resume_formatter import *
from utils.resume_utils import _register_palatino_fonts
import utils.resume_utils as ru
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import json
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_register_palatino_fonts()

def create_resume_docx_for_new_hires(template_bytes, data):
    if isinstance(data, str):
        data = json.loads(data)
    doc = Document()
    # PAGE SETUP
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.1)
    section = doc.sections[0]

    # FONT SETUP
    style = doc.styles["Normal"]
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1
    style.font.name = "Palatino Linotype"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"),"Palatino Linotype")
    style._element.rPr.rFonts.set(qn("w:ascii"),"Palatino Linotype")
    style._element.rPr.rFonts.set(qn("w:hAnsi"),"Palatino Linotype")

    # HELPERS
    def remove_table_borders(table):
        tbl = table._tbl
        tblPr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for border_name in ["top","left","bottom","right","insideH","insideV",]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "nil")
            borders.append(border)
        tblPr.append(borders)

    def add_page_field(run, field_name):
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = field_name
        fld_separate = OxmlElement("w:fldChar")
        fld_separate.set(qn("w:fldCharType"), "separate")
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_separate)
        run._r.append(fld_end)

    '''def add_horizontal_line():
        """
        Adds a real Word horizontal line using paragraph borders.
        Replaces underscore-based separators.
            """
        p = doc.add_paragraph()
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p'''

    def add_section(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Palatino Linotype"
        run.font.size = Pt(11)
        pPr = p._element.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), "000000")
        pBdr.append(bottom)
        pPr.append(pBdr)

    def add_bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        run = p.add_run(str(text))
        run.font.name = "Palatino Linotype"

    def add_bold_line(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        r1 = p.add_run(label)
        r1.bold = True
        r1.font.name = "Palatino Linotype"
        r1.font.size = Pt(10)
        r2 = p.add_run(str(value))
        r2.font.name = "Palatino Linotype"
        r2.font.size = Pt(10)

    # LOGO
    logo_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),"assets","kumaran-logo.png")
    if os.path.exists(logo_path):
        header = doc.sections[0].header
        p = header.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        run = p.add_run()
        run.add_picture(logo_path,width=Inches(1.2))

    # NAME
    name = data.get("name", "")
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
    run = p.add_run(name)
    run.bold = True
    run.font.name = "Palatino Linotype"
    run.font.size = Pt(16)

    # PROFESSIONAL SUMMARY
    summary = data.get("professional_summary", [])
    if summary:
        add_section("Professional Summary")
        for item in summary:
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
            r = p.add_run(str(item))
            r.font.name = "Palatino Linotype"

    # EMPLOYMENT HISTORY
    employment = data.get("employment_history",[])
    if employment:
        add_section("Employment History")
        for emp in employment:
            designation = str(emp.get("designation", "")).strip()
            company = str(emp.get("company", "")).strip()
            start = str(emp.get("start_date", "")).strip()
            end = str(emp.get("end_date") or "Present").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            bullet = p.add_run("• ")
            bullet.font.name = "Palatino Linotype"
            r1 = p.add_run(designation)
            r1.bold = True
            r1.font.name = "Palatino Linotype"
            r1.font.size = Pt(10)
            r2 = p.add_run(f" - {company}")
            r2.font.name = "Palatino Linotype"
            r2.font.size = Pt(10)
            r3 = p.add_run(f" - {start} to {end}")
            r3.bold = True
            r3.font.name = "Palatino Linotype"
            r3.font.size = Pt(10)

    # TECHNICAL SKILLS
    technical_skills = data.get("technical_skills",{})
    if technical_skills:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        add_section("Technical Skills")
        table = doc.add_table(rows=0,cols=3)
        table.autofit = False
        remove_table_borders(table)
        # Expand "other" section
        if "other" in technical_skills:
            other = technical_skills.get("other", {})
            if isinstance(other, dict):
                for key, value in other.items():
                    if value:
                        technical_skills[key] = value
            technical_skills.pop("other", None)
        for k, v in technical_skills.items():
            if isinstance(v, list):
                value = ", ".join(v)
            elif isinstance(v, dict):
                continue
            else:
                value = str(v)
            if not value.strip():
                continue
            label = (k.replace("_", " ").title())
            row = table.add_row()
            row.cells[0].width = Inches(2.2)
            row.cells[1].width = Inches(0.2)
            row.cells[2].width = Inches(4.6)
            p1 = row.cells[0].paragraphs[0]
            p2 = row.cells[1].paragraphs[0]
            p3 = row.cells[2].paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            p3.paragraph_format.space_before = Pt(0)
            p3.paragraph_format.space_after = Pt(0)
            # Bullet like PDF
            p1.add_run("• ")
            r1 = p1.add_run(label)
            r1.bold = True
            r1.font.name = "Palatino Linotype"
            r2 = p2.add_run(":")
            r2.font.name = "Palatino Linotype"
            r3 = p3.add_run(value)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
      
    # PROFESSIONAL EXPERIENCE
    experience = data.get("professional_experience",[])
    if experience:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        add_section("Professional Experience")
        for exp in experience:
            company = str(exp.get("company", "")).strip()
            start = str(exp.get("start_date", "")).strip()
            end = str(exp.get("end_date") or "Present").strip()
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(3)
            table = doc.add_table(rows=1,cols=2)
            table.autofit = False
            remove_table_borders(table)
            left = table.cell(0, 0)
            right = table.cell(0, 1)
            left.width = Inches(5.0)
            right.width = Inches(2.2)
            lp = left.paragraphs[0]
            rp = right.paragraphs[0]
            r1 = lp.add_run(company)
            r1.bold = True
            r1.font.name = "Palatino Linotype"
            rp.alignment = (WD_PARAGRAPH_ALIGNMENT.RIGHT)
            r2 = rp.add_run(f"{start} - {end}")
            r2.bold = True
            r2.font.name = "Palatino Linotype"
            line = doc.add_paragraph()
            line.paragraph_format.space_before = Pt(0)
            line.paragraph_format.space_after = Pt(3)
            line.paragraph_format.line_spacing = 0
            pPr = line._element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "000000")
            pBdr.append(bottom)
            pPr.append(pBdr)
            r2.bold = True
            r2.font.name = "Palatino Linotype"
            projects = exp.get("projects",[])
            projects = exp.get("projects",[])
            for proj in projects:
                role = proj.get("role","")
                project_name = proj.get("project_name","")
                client_name = proj.get("client_name","")
                domain = proj.get("domain","")
                desc = proj.get("project_description",[])
                roles = proj.get("roles_responsibilities",[])
                technologies = proj.get("technologies",[])
                if role:
                    add_bold_line("Role: ",role)
                if project_name:
                    add_bold_line("Project Name: ",project_name)
                if client_name:
                    add_bold_line("Client Name: ",client_name)
                if domain:
                    add_bold_line("Domain: ",domain)
                if technologies:
                    add_bold_line("Technologies and tools used: ",", ".join(technologies))
                if desc:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1
                    p.alignment = (WD_PARAGRAPH_ALIGNMENT.JUSTIFY)
                    r = p.add_run("Project Description: ")
                    r.bold = True
                    r.font.name = "Palatino Linotype"
                    if isinstance(desc, list):
                        desc = " ".join(desc)
                    r2 = p.add_run(desc)
                    r2.font.name = "Palatino Linotype"
                if roles:
                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.space_after = Pt(3)
                    heading = doc.add_paragraph()
                    heading.paragraph_format.space_before = Pt(0)
                    heading.paragraph_format.space_after = Pt(2)
                    heading.paragraph_format.line_spacing = 1
                    rh = heading.add_run("Roles & Responsibilities:")
                    rh.bold = True
                    rh.font.name = "Palatino Linotype"
                    rh.font.size = Pt(10)
                    for item in roles[:6]:
                        add_bullet(item)
                    spacer = doc.add_paragraph()
                    spacer.paragraph_format.space_after = Pt(3)

    # EDUCATION
    education = data.get("education",[])
    if education:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        add_section("Education")
        def get_start_year(e):
            years = str(e.get("year", "")).split("-")
            return (years[0].strip()if years else "")
        education_list = sorted(education,key=lambda x: get_start_year(x),reverse=True)
        for edu in education_list:
            degree = str(edu.get("degree", "")).strip()
            stream = str(edu.get("stream", "")).strip()
            institution = str(edu.get("institution","")).strip()
            score = str(edu.get("percentage","")).strip()
            year = str(edu.get("year","")).strip()
            parts = []
            if degree:
                if stream:
                    parts.append(f"{degree} ({stream})")
                else:
                    parts.append(degree)
            if institution:
                parts.append(institution)
            if score:parts.append(score)
            if year:
                parts.append(year)
            add_bullet(" – ".join(parts))

    # FOOTER
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = (WD_PARAGRAPH_ALIGNMENT.LEFT)
    run = footer_para.add_run(f"Profile of {name} (Page ")
    run.font.name = "Palatino Linotype"
    run.font.size = Pt(9)
    page_run = footer_para.add_run()
    page_run.font.name = "Palatino Linotype"
    add_page_field(page_run,"PAGE")
    run = footer_para.add_run(" of ")
    run.font.name = "Palatino Linotype"
    total_run = footer_para.add_run()
    total_run.font.name = "Palatino Linotype"
    add_page_field(total_run,"NUMPAGES")
    run = footer_para.add_run(")")
    run.font.name = "Palatino Linotype"

    # SAVE
    output_docx_buffer = BytesIO()
    doc.save(output_docx_buffer)
    output_docx_buffer.seek(0)
    return output_docx_buffer

def create_resume_pdf_for_new_hires(template_bytes, data):
    if isinstance(data, str):
        data = json.loads(data)
    ru.initialize_globals(template_bytes, data)
    ru.new_page()

    # SUMMARY
    if data.get("professional_summary"):
        ru.section("Professional Summary")
        for item in data["professional_summary"]:
            # bullet symbol
            indent = ru.BASE_BULLET_INDENT
            bullet_x = ru.LEFT_MARGIN + indent
            text_x = bullet_x + 10
            max_width = ru.RIGHT_MARGIN - text_x
            ru.resume_canvas.setFont(ru.DEFAULT_FONT, 10)
            ru.resume_canvas.drawString(bullet_x, ru.content_position, u"\u2022")
            # justified text
            ru.draw_text_justified(item, text_x, max_width)

    if data.get("employment_history"):
        ru.section("Employment History")
        for emp in data["employment_history"]:
            #  add spacing BEFORE next line (except first)
            if ru.content_position  != ru.TOP:
                ru.content_position -= 6   #  visible gap between ru.bullets
            designation = str(emp.get("designation", "")).strip()
            company = str(emp.get("company", "")).strip()
            start = str(emp.get("start_date", "")).strip()
            end = str(emp.get("end_date") or "Present").strip()
            indent = ru.BASE_BULLET_INDENT
            bullet_x = ru.LEFT_MARGIN + indent
            text_x = bullet_x + 12
            date_text = f"{start} to {end}"
            ru.check_space()
            ru.resume_canvas.setFont("Palatino", 10)
            ru.resume_canvas.drawString(bullet_x, ru.content_position, u"\u2022")
            ru.resume_canvas.setFont("Palatino-Bold", 10)
            ru.resume_canvas.drawString(text_x, ru.content_position, designation)
            role_width = ru.resume_canvas.stringWidth(designation, "Palatino-Bold", 10)
            company_x = text_x + role_width + 5
            ru.resume_canvas.setFont("Palatino", 10)
            ru.resume_canvas.drawString(company_x, ru.content_position, f" - {company}")
            company_width = ru.resume_canvas.stringWidth(f" - {company}", "Palatino", 10)
            date_x = company_x + company_width + 5
            ru.resume_canvas.setFont("Palatino-Bold", 10)
            ru.resume_canvas.drawString(date_x, ru.content_position, f" – {date_text}")
            #  normal line move
            ru.content_position -= ru.LINE_HEIGHT
    
    # TECHNICAL SKILLS
    if data.get("technical_skills"):
        ru.section("Technical Skills")
        SKILL_LABEL_WIDTH = 170   # adjust for perfect alignment
        for k, v in data["technical_skills"].items():
            label = k.replace('_', ' ').title()
                # Convert value
            value = ", ".join(v) if isinstance(v, list) else str(v)
                # IMPORTANT: skip empty values
            if not value or value.strip() == "":
                continue
            indent = ru.BASE_BULLET_INDENT
            bullet_x = ru.LEFT_MARGIN + indent
            label_x = bullet_x + 12
            colon_x = label_x + SKILL_LABEL_WIDTH
            value_x = colon_x + 12
            ru.check_space()
            # Bullet
            ru.resume_canvas.setFont(ru.DEFAULT_FONT, 10)
            ru.resume_canvas.drawString(bullet_x, ru.content_position, u"\u2022")
            # Label (bold like image)
            ru.resume_canvas.setFont(ru.DEFAULT_BOLD, 10)
            ru.resume_canvas.drawString(label_x, ru.content_position, label)
            # Colon aligned
            ru.resume_canvas.setFont(ru.DEFAULT_BOLD, 10)
            ru.resume_canvas.drawString(colon_x, ru.content_position, ":")
            # Value (normal font, wrapped properly)
            ru.draw_text(value, x=value_x)
            ru.content_position -= ru.LINE_HEIGHT
    
    if data.get("professional_experience"):
        ru.section("Professional Experience")
        ru.content_position -= ru.LINE_HEIGHT * 2
        for exp in data["professional_experience"]:
            # check_space(160)
            ru.check_space(ru.LINE_HEIGHT * 5)
            company = str(exp.get("company", "")).strip()
            start = str(exp.get("start_date", "")).strip()
            end = str(exp.get("end_date") or "Present").strip()
            date_str = f"{start} - {end}"
            projects = exp.get("projects") or []
            # COMPANY + DATE
            if company:
                ru.resume_canvas.setFont(ru.DEFAULT_BOLD, 10)
                ru.resume_canvas.drawString(ru.LEFT_MARGIN, ru.content_position, company)
                ru.resume_canvas.drawRightString(ru.RIGHT_MARGIN, ru.content_position, date_str)
                ru.content_position -= 4
                ru.resume_canvas.line(ru.LEFT_MARGIN, ru.content_position, ru.RIGHT_MARGIN,ru.content_position)
                ru.content_position -= ru.LINE_HEIGHT

            # LOOP PROJECTS
            for proj in projects:
                ru.check_space(140)
                project_name = proj.get("project_name", "")
                client_name = proj.get("client_name", "")
                domain = proj.get("domain", "")
                role = proj.get("role", "")   # FIXED LOCATION
                desc = proj.get("project_description", [])
                roles = proj.get("roles_responsibilities", [])
                tech_list = proj.get("technologies", [])
                # ROLE (FIRST)
                if role:
                    ru.draw_text(f"Role: {role}", bold=True)
                    ru.content_position -= ru.LINE_HEIGHT
                # PROJECT NAME
                if project_name:
                    ru.draw_text(f"Project Name: {project_name}", bold=True)
                    ru.content_position -= ru.LINE_HEIGHT
                # CLIENT NAME
                if client_name:
                    ru.draw_text(f"Client Name: {client_name}",bold=True)
                    ru.content_position -= ru.LINE_HEIGHT
                # DOMAIN
                if domain:
                    ru.draw_text(f"Domain: {domain}",bold=True)
                    ru.content_position -= ru.LINE_HEIGHT
                # TECHNOLOGIES
                if tech_list:
                    tech_text = ", ".join(tech_list)
                    ru.resume_canvas.setFont(ru.DEFAULT_BOLD, 10)
                    ru.resume_canvas.drawString(ru.LEFT_MARGIN, ru.content_position, "Technologies and tools used:")
                    text_x = ru.LEFT_MARGIN + ru.resume_canvas.stringWidth("Technologies and tools used: ","Palatino-Bold",10)
                    ru.resume_canvas.setFont(ru.DEFAULT_FONT, 10)
                    ru.draw_text(tech_text, x=text_x)
                    ru.content_position -= ru.LINE_HEIGHT
                # PROJECT DESCRIPTION (PARAGRAPH)
                if desc:
                    ru.draw_text("Project Description:", bold=True)
                    if isinstance(desc, list):
                        desc = " ".join(desc)   # convert to paragraph
                    ru.draw_text(desc)
                    ru.content_position -= ru.LINE_HEIGHT
                # ROLES & RESPONSIBILITIES (BULLETS)       
                    if roles:
                        ru.draw_text("Roles & Responsibilities:", bold=True)
                        roles = roles[:6]   # limit
                        for r in roles:
                            ru.check_space(ru.LINE_HEIGHT * 2)
                            ru.bullet(r, level=1)
                        ru.content_position -= ru.LINE_HEIGHT
                    # GAP BETWEEN PROJECTS
                    ru.content_position -= ru.LINE_HEIGHT
    if data.get("education"):
        ru.section("Education")
        # Sort by start year (descending → latest first)
        def get_start_year(e):
            years = str(e.get("year", "")).split("-")
            return years[0].strip() if years else ""
        education_list = sorted(data["education"],key=lambda x: get_start_year(x),reverse=True)
        for edu in education_list:
            ru.check_space(ru.LINE_HEIGHT * 2)
            degree = str(edu.get("degree", "")).strip()
            stream = str(edu.get("stream", "")).strip()
            institution = str(edu.get("institution", "")).strip()
            score = str(edu.get("percentage", "")).strip()
            year = str(edu.get("year", "")).strip()
            # Build text properly
            parts = []
            if degree:
                if stream:
                    parts.append(f"{degree} ({stream})")
                else:
                    parts.append(degree)
            if institution:
                parts.append(institution)
            if score:
                parts.append(score)
            if year:
                parts.append(year)
            full_text = " – ".join(parts)
            # Bullet render
            ru.bullet(full_text, level=0)
            ru.content_position -= 4
    if ru.resume_canvas:
        ru.resume_canvas.save()
        ru.buffer.seek(0)
        ru.pages.append(ru.buffer)
    total_pages = len(ru.pages)
    writer = PdfWriter()
    for i, pg in enumerate(ru.pages, start=1):
        content = PdfReader(pg)
        footer_buf = BytesIO()
        fc = canvas.Canvas(footer_buf, pagesize=letter)
        fc.setFont("Palatino-Italic", 8)
        fc.drawString(ru.LEFT_MARGIN, ru.FOOTER_Y1,f"Profile of {ru.name} (Page {i} of {total_pages})")
        fc.save()
        footer_buf.seek(0)
        footer_pdf = PdfReader(footer_buf)
        base_page = copy.copy(ru.template_pdf.pages[0])
        base_page.merge_page(content.pages[0])
        # THIRD → footer
        base_page.merge_page(footer_pdf.pages[0])
        writer.add_page(base_page)
    output = BytesIO()
    writer.write(output)
    output.seek(0)
    return output





