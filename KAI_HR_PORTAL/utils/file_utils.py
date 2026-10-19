import tempfile
from pathlib import Path
import pdfplumber
from docx import Document
import os
from logging_config.logger import logger
try:
    from pypdf import PdfReader  # pypdf (newer)
except ImportError:
    try:
        from PyPDF2 import PdfReader  # type: ignore  # PyPDF2 (fallback, optional)
    except ImportError:
        PdfReader = None

def extract_pdf_text(file_storage):
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        text_parts = []
        urls_found = []
        
        # Extract visible text using pdfplumber
        with pdfplumber.open(tmp_path) as pdf:
            for i, page in enumerate(pdf.pages):
                if page_text := page.extract_text():
                    text_parts.append(page_text)
        
        # Extract hidden hyperlinks and annotations using pypdf/PyPDF2
        if PdfReader:
            try:
                reader = PdfReader(tmp_path)
                for page_num, page in enumerate(reader.pages):
                    # Extract annotations (links)
                    if "/Annots" in page:
                        annotations = page["/Annots"]
                        for annotation in annotations:
                            annot_obj = annotation.get_object()
                            
                            # Extract URL from link annotations
                            if annot_obj.get("/Subtype") == "/Link":
                                if "/A" in annot_obj:
                                    action = annot_obj["/A"]
                                    if "/URI" in action:
                                        url = action["/URI"]
                                        urls_found.append(url)
                            
                            # Extract text from annotations (comments, notes)
                            if "/Contents" in annot_obj:
                                content = annot_obj["/Contents"]
                                if content and content.strip():
                                    urls_found.append(f"Annotation: {content}")
            except Exception as e:
                # If link extraction fails, continue with just visible text
                logger.info(f"Warning: Could not extract links: {e}")
        
        # Combine visible text with extracted URLs
        full_text = "\n".join(text_parts)
        if urls_found:
            full_text += "\n\n--- Extracted Links and Annotations ---\n"
            full_text += "\n".join(urls_found)
        
        return full_text
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def extract_doc_text(file_storage):
    """
    Extract text and hyperlinks from legacy .doc files (Word 97-2003 format).
    Uses Word COM automation for full extraction on Windows.
    Falls back to textract (text only) if Word COM is unavailable.
    """
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        # Method 1: Try Word COM automation (Windows only, extracts text + hyperlinks)
        if os.name == 'nt':  # Windows
            try:
                import win32com.client
                
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(tmp_path)
                
                # Extract headers from all sections
                header_text = []
                try:
                    for section in doc.Sections:
                        # Extract headers (primary, first page, even pages)
                        for header in section.Headers:
                            if header.Range.Text.strip():
                                header_text.append(header.Range.Text.strip())
                        # Extract footers
                        for footer in section.Footers:
                            if footer.Range.Text.strip():
                                header_text.append(footer.Range.Text.strip())
                except Exception as e:
                    logger.info(f"Warning: Could not extract headers from .doc: {e}")
                
                # Extract visible text from body
                text_content = doc.Content.Text
                
                # Extract hyperlinks
                urls_found = []
                try:
                    for hyperlink in doc.Hyperlinks:
                        if hyperlink.Address:
                            urls_found.append(hyperlink.Address)
                except Exception as e:
                    logger.info(f"Warning: Could not extract hyperlinks from .doc: {e}")
                
                # Close Word
                doc.Close(False)
                word.Quit()
                
                # Combine header text with body text and hyperlinks
                full_text = ""
                if header_text:
                    full_text += "--- Headers ---\n"
                    full_text += "\n".join(header_text)
                    full_text += "\n\n--- Document Body ---\n"
                full_text += text_content
                if urls_found:
                    full_text += "\n\n--- Extracted Links ---\n"
                    full_text += "\n".join(urls_found)
                
                return full_text
                
            except Exception as e:
                logger.info(f"Warning: Word COM automation failed: {e}. Falling back to textract...")
        
        # Method 2: Fallback to textract (text only, no hyperlinks)
        try:
            import textract
            text = textract.process(tmp_path).decode('utf-8')
            return text
        except ImportError:
            raise Exception(
                "Unable to process .doc files because the required document processing dependencies are not installed."
            )
        except Exception as e:
            raise Exception(f"Failed to process the uploaded .doc file. Please convert the document to .docx format and try again.{str(e)}")
    
    finally:
        Path(tmp_path).unlink(missing_ok=True)


def extract_docx_text(file_storage):
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        file_storage.save(tmp.name)
        tmp_path = tmp.name

    try:
        doc = Document(tmp_path)
        text_parts = []
        urls_found = set()  # Use set to avoid duplicates
        
        # Extract headers from all sections
        header_text = []
        seen_text = set()  # Track unique text to avoid duplicates
        for section_idx, section in enumerate(doc.sections):
            # Extract from header (primary header)
            header = section.header
            
            # Method 1: Extract text from header paragraphs
            para_count = len(header.paragraphs)
            table_count = len(header.tables)
            logger.info(f"DEBUG: Section {section_idx} header has {para_count} paragraphs, {table_count} tables")
            
            for para_idx, paragraph in enumerate(header.paragraphs):
                text = paragraph.text.strip()
                logger.info(f"DEBUG: Header paragraph {para_idx}: '{text[:50] if text else '(empty)'}'")
                if text and text not in seen_text:
                    header_text.append(text)
                    seen_text.add(text)
            
            # Method 2: Extract text from tables in header
            for table_idx, table in enumerate(header.tables):
                logger.info(f"DEBUG: Header table {table_idx} has {len(table.rows)} rows")
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            text = paragraph.text.strip()
                            if text and text not in seen_text:
                                logger.info(f"DEBUG: Found text in table cell: '{text[:50]}'")
                                header_text.append(text)
                                seen_text.add(text)
            
            # Method 3: Extract text from text boxes and shapes in header (using XML)
            try:
                from lxml import etree

                # Define namespace
                nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

                # Find all text box content elements (w:txbxContent)
                txbx_elements = header._element.findall('.//w:txbxContent', nsmap)

                logger.info(f"DEBUG: Found {len(txbx_elements)} text boxes in header using findall")

                for txbx_idx, txbx in enumerate(txbx_elements):
                    # Extract all text elements within the text box
                    text_elements = txbx.findall('.//w:t', nsmap)
                    for t_elem in text_elements:
                        text = t_elem.text
                        if text and text.strip():
                            text_clean = text.strip()
                            if text_clean not in seen_text:
                                logger.info(f"DEBUG: Found text in text box {txbx_idx}: '{text_clean[:50]}'")
                                header_text.append(text_clean)
                                seen_text.add(text_clean)

                # Alternative method: Get all text from the entire header element
                if not txbx_elements:
                    logger.info("DEBUG: No text boxes found, trying alternative extraction...")
                    # Get all w:t elements (text runs) from header
                    all_text_elements = header._element.findall('.//w:t', nsmap)
                    logger.info(f"DEBUG: Found {len(all_text_elements)} text elements in header")

                    for t_elem in all_text_elements:
                        text = t_elem.text
                        if text and text.strip():
                            text_clean = text.strip()
                            if text_clean not in seen_text:
                                logger.info(f"DEBUG: Found text element: '{text_clean[:50]}'")
                                header_text.append(text_clean)
                                seen_text.add(text_clean)

            except Exception as e:
                logger.info(f"DEBUG: Could not extract text boxes from header: {e}")
                import traceback
                traceback.print_exc()
        
        # Extract visible text and hyperlinks from body
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
            
            # Extract hyperlinks from paragraph
            for hyperlink in paragraph._element.xpath('.//w:hyperlink'):
                # Get the relationship ID
                rel_id = hyperlink.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                if rel_id:
                    try:
                        # Get the actual URL from document relationships
                        url = doc.part.rels[rel_id].target_ref
                        if url:
                            urls_found.add(url)
                    except (KeyError, AttributeError):
                        pass
        
        # Also check tables for hyperlinks
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for hyperlink in paragraph._element.xpath('.//w:hyperlink'):
                            rel_id = hyperlink.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                            if rel_id:
                                try:
                                    url = doc.part.rels[rel_id].target_ref
                                    if url:
                                        urls_found.add(url)
                                except (KeyError, AttributeError):
                                    pass
        
        # Combine header text with body text and URLs
        full_text = ""
        if header_text:
            logger.info(f"DEBUG: Extracted {len(header_text)} header lines")
            full_text += "--- Headers ---\n"
            full_text += "\n".join(header_text)
            full_text += "\n\n--- Document Body ---\n"
        else:
            logger.info("DEBUG: No header text found")
        full_text += "\n".join(text_parts)
        if urls_found:
            full_text += "\n\n--- Extracted Links ---\n"
            full_text += "\n".join(sorted(urls_found))
        
        return full_text
    finally:
        Path(tmp_path).unlink(missing_ok=True)