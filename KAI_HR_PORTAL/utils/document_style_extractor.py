import io
from io import BytesIO
import json
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from collections import defaultdict
import tempfile

def get_resolved_font(run, paragraph):
    if run.font.name:
        return run.font.name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is not None:
        theme = rFonts.get(qn('w:asciiTheme')) or rFonts.get(qn('w:minorHAnsi'))
        if theme: return "Calibri"
    return paragraph.style.font.name or "Calibri"

def get_run_color(run):
    if run.font.color and run.font.color.rgb:
        return str(run.font.color.rgb)
    return "000000"

def aggregate_issue(storage, issue_type, detected_val, text):
    key = (issue_type, str(detected_val))
    if key not in storage:
        storage[key] = []
    storage[key].append(text)


def analyze_letter_format(doc, response):
    grouped_issues = {} 
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text: continue
        align = para.alignment if para.alignment is not None else para.style.paragraph_format.alignment
        if align != WD_ALIGN_PARAGRAPH.JUSTIFY:
            aggregate_issue(grouped_issues, "alignment_issue", str(align), text)
        sa = para.paragraph_format.space_after
        if sa:
            line_val = round(sa.pt / 12.0, 2)
            if not (1.0 <= line_val <= 1.5):
                aggregate_issue(grouped_issues, "Paragraph_Spacing_issue", f"spacing_{line_val}", text)
        else:
            aggregate_issue(grouped_issues, "Paragraph_Spacing_issue", "no_spacing", text)
        for run in para.runs:
            r_text = run.text.strip()
            if not r_text: continue
            f_name = get_resolved_font(run, para)
            f_size = run.font.size.pt if run.font.size else 11
            if f_name.lower() != "calibri":
                aggregate_issue(grouped_issues, "font_family", f_name, r_text)
            if not (11 <= f_size <= 12):
                aggregate_issue(grouped_issues, "font_size", f_size, r_text)

    for (i_type, i_val), texts in grouped_issues.items():
        letter_issue = {
            "type": i_type,
            f"detected_{i_type}": i_val, 
            "text_with_issue": " ".join(texts)
            }
        if i_type == "font_size":
            letter_issue['expected_font_size'] = "11 - 12"
        elif i_type == "font_family":
            letter_issue['expected_font_family'] = "calibri"
        response["issues"].append(letter_issue)
    return response

def analyze_resume_format(doc, response):
    grouped_issues = {}
    page_regex = r"([A-Za-z\.\s]+)\s?\(Page \d+ of \d+\)"
    email_regex = r"[\w\.-]+@[\w\.-]+\.\w+"
    link_regex = r"(https?://(www\.)?linkedin\.com)"

    # 1. Header & Footer Analysis 
    for i, section in enumerate(doc.sections):
        headers = [section.header, section.first_page_header, section.even_page_header]
        footers = [section.footer, section.first_page_footer, section.even_page_footer]

        for h_idx, header in enumerate(headers):
            h_paras = header.paragraphs
            h_text = " ".join([p.text for p in h_paras]).strip()
            if h_text:
                for para in h_paras:
                    for run in para.runs:
                        if not run.text.strip(): continue
                        f_name = get_resolved_font(run, para)
                        f_size = run.font.size.pt if run.font.size else 0
                        if f_name.lower() != "palatino linotype":
                            aggregate_issue(grouped_issues, "header_font_family", f_name, run.text)
                        if f_size != 16:
                            aggregate_issue(grouped_issues, "header_font_size", f_size, run.text)
                        if get_run_color(run) != "000000":
                            aggregate_issue(grouped_issues, "header_font_color", f"#{get_run_color(run)}", run.text)

        for f_idx, footer in enumerate(footers):
            f_paras = footer.paragraphs
            f_text = " ".join([p.text for p in f_paras]).strip()
            if f_text:
                # Content checks (Regex)
                if not re.search(page_regex, f_text):
                    response["issues"].append({"type": "footer_content", "issue": "Missing 'Name (Page X of Y)' format"})
                if not re.search(link_regex, f_text):
                    response["issues"].append({"type": "footer_content", "issue": "Missing LinkedIn URL"})
                if not re.search(email_regex, f_text):
                    response["issues"].append({"type": "footer_content", "issue": "Missing Email"})

                for para in f_paras:
                    for run in para.runs:
                        if not run.text.strip(): continue
                        f_name = get_resolved_font(run, para)
                        f_size = run.font.size.pt if run.font.size else 0
                        if f_name.lower() != "palatino linotype":
                            aggregate_issue(grouped_issues, "footer_font_family", f_name, run.text)
                        if f_size != 8:
                            aggregate_issue(grouped_issues, "footer_font_size", f_size, run.text)
                        if not run.font.italic:
                            aggregate_issue(grouped_issues, "footer_font_style", "not_italic", run.text)
                        if get_run_color(run) != "000000":
                            aggregate_issue(grouped_issues, "footer_font_color",f"#{get_run_color(run)}", run.text)

    # 2. Body Analysis
    for para in doc.paragraphs:
        if not para.text.strip(): continue
        for run in para.runs:
            r_text = run.text.strip()
            if not r_text: continue
            f_name = get_resolved_font(run, para)
            f_size = run.font.size.pt if run.font.size else 10
            f_color = get_run_color(run)
            if f_name.lower() != "palatino linotype":
                aggregate_issue(grouped_issues, "font_family", f_name, r_text)
            if f_size != 10:
                aggregate_issue(grouped_issues, "font_size", f_size, r_text)
            if f_color != "000000":
                aggregate_issue(grouped_issues, "font_color_issue", f"#{f_color}", r_text)

    for (i_type, i_val), texts in grouped_issues.items():
        resume_issue = {
            "type": i_type,
            f"detected_{i_type}": i_val, 
            "text_with_issue": " ".join(texts)
            }
        if i_type == "font_size":
            resume_issue['expected_font_size'] = "10"

        elif i_type == "font_family":
            resume_issue['expected_font_family'] = "palatino linotype"

        elif i_type == "font_color_issue":
            resume_issue['expected_font_color'] = "#000000"
        
        elif i_type == "footer_font_family":
            resume_issue['expected_footer_font_family'] = "palatino linotype"
        
        elif i_type == "footer_font_color":
            resume_issue["expected_footer_font_color"] = "#000000"
        
        elif i_type == "footer_font_style":
            resume_issue["expected_footer_font_style"] = "italic"

        elif i_type == "header_font_family":
            resume_issue["expected_header_font_family"] = "palatino linotype" 

        elif i_type == "header_font_size":
            resume_issue['expected_header_font_size'] = 16
        
        elif i_type == "header_font_color":
            resume_issue['expected_header_font_color'] = "#000000"
        
        response["issues"].append(resume_issue)
    
    
    return response

def analyze_document(file_bytes, document_type: str, document_name="document.docx") -> dict:
    response = {
        "document_name": document_name,
        "document_type": document_type,
        "is_corupted": False,
        "issues": [],
        "is_pdf":False
    }
    # file_bytes = file_bytes.read()

    # if file_bytes.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
    #     response["issues"].append({"type": "depricated document format", "issue": "Please use .docx"})
    #     return response
    # if not file_bytes.startswith(b'PK\x03\x04'):
    #     response["is_corupted"] = True
    #     return response
    try:
        doc = Document(file_bytes)
    except:
        response["is_corupted"] = True
        return response

    if document_type.lower() == "letter":
        return analyze_letter_format(doc, response)
    elif document_type.lower() == "resume":
        return analyze_resume_format(doc, response)
    else:
        return {"error":"unsupported document type ,Supported formats are letter,resume"}