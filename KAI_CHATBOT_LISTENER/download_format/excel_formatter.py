import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, Border, Side
from io import BytesIO 
from flask import make_response,jsonify


def excel_construction(applied_filters,df_data,type,platform_language):
    try:
        print(type)
        report_type_translations_french = {
            'performance': 'performance',
            'feedback': 'retours',
            'tokenutilized': 'tokens_utilisés',
            'usage': 'utilisation'
        }
   
        if platform_language=="en":
            report_types = [rt for rt in type]
        elif platform_language=="fr":
            report_types = [report_type_translations_french.get(rt, rt) for rt in type]
            
        report_header = ",".join(report_types)
        print(report_header)
        df_data.index = df_data.index + 1
        filter_data = [{'Filter': key, 'Value': applied_filters.get(key, 'ALL')} for key in ['usernames', 'response_time', 'response_types','ratings', 'year', 'week', 'month', 'start_date', 'end_date', 'language_type','model']]
        print(filter_data)
        if platform_language!='en':
            filter_data =language_based_headers(filter_data=filter_data)
        df_filters = pd.DataFrame(filter_data)
        output = BytesIO()
        df_filters.to_excel(output, index=False, startrow=2)
        with pd.ExcelWriter(output, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            df_data.to_excel(writer, index=True, startrow=15, startcol=0)
        output.seek(0)
        wb = load_workbook(output)
        ws = wb.active
        thin_border = Border(left=Side(border_style="thin"), 
                    right=Side(border_style="thin"), 
                    top=Side(border_style="thin"), 
                    bottom=Side(border_style="thin"))
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        ws['A1'] = f"{report_header}"
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A1'].font = Font(size=14, bold=True)
        if 'group_by' in applied_filters and applied_filters['group_by'] and not df_data.empty:
            ws.merge_cells(start_row=15, start_column=1, end_row=15, end_column=2)
            ws['A15'] = "Grouper par" if platform_language == 'fr' else "Group By"
            ws['A15'].alignment = Alignment(horizontal='left', vertical='center')
            ws['A15'].font = Font(size=14, bold=True)
        ws.merge_cells(start_row=15, start_column=1, end_row=15, end_column=2)
        ws['A16'] = "Sortie du rapport" if platform_language == 'fr' else "Report Output"
        ws['A16'].alignment = Alignment(horizontal='center', vertical='center')
        ws['A16'].font = Font(size=14, bold=True)
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = thin_border
        output.seek(0)
        wb.save(output)
        output.seek(0)
        response = make_response(output.getvalue())
        response.headers['Content-Disposition'] = 'attachment; filename=Reportdata.xlsx'
        response.headers['Content-type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        return response
    except Exception as e:
        return make_response(jsonify(error=f"Error in download excel {e}"),500)
    
def language_based_headers(filter_data):
    df = pd.DataFrame(filter_data)
    french_mapping = {'usernames':'noms_utilisateurs', 'response_time':'temps_de_réponse', 'response_types':'types_de_réponse', 'ratings':'évaluations', 'year':'année', 'week':'semaine', 'month':'mois', 'start_date':'date_de_début', 'end_date':'date_de_fin', 'language_type':'type_de_langue', 'model':'modèle'}
    df['Filter'] = df['Filter'].map(french_mapping).fillna(df['Filter'])
    df = df.rename(columns={'Filter': 'Filtre', 'Value': 'Valeur'})
    filter = df.to_dict(orient='records')
    return filter